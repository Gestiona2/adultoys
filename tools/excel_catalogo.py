#!/usr/bin/env python3
"""Puente entre el catalogo (JSON) y la planilla del dueno (Excel).

    python3 tools/excel_catalogo.py generar
    python3 tools/excel_catalogo.py revisar Catalogo-Adultoyss.xlsx
    python3 tools/excel_catalogo.py aplicar Catalogo-Adultoyss.xlsx

El dueno edita la planilla y la devuelve. `revisar` la compara contra el JSON
con las mismas reglas del validador y dice, producto por producto, que quedo
mal y como se escribe bien. `aplicar` solo escribe si no hay ningun error.

El descuento NO se lee de la planilla: se calcula de los dos precios con la
misma cuenta del validador, asi siempre cuadra. En la planilla hay una formula
que lo muestra en vivo para orientar al dueno.

La categoria y la subcategoria se muestran con su nombre legible. Al aplicar
se vuelven a traducir a los ids internos del catalogo.
"""
import json
import re
import sys
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from validar import validar, cargar_categorias, descuento_esperado

RAIZ = Path(__file__).resolve().parent.parent
RUTA_JSON = RAIZ / "src" / "datos" / "productos.json"
RUTA_CATEGORIAS = RAIZ / "src" / "datos" / "categorias.json"
RUTA_XLSX = RAIZ / "Catalogo-Adultoyss.xlsx"

FUENTE = Font(name="Arial", size=11)
FUENTE_TITULO = Font(name="Arial", size=11, bold=True, color="FFFFFF")
FONDO_TITULO = PatternFill("solid", fgColor="3A0A2A")
AMARILLO = PatternFill("solid", fgColor="FFF2CC")   # celdas que edita el dueno
GRIS = PatternFill("solid", fgColor="EFEFEF")       # no tocar
CALCULO = PatternFill("solid", fgColor="D9EAD3")    # se calcula solo
VERDE = PatternFill("solid", fgColor="E2F0D9")      # estado calculado

# slug, titulo, marca, presentacion, categoria, subcategoria, precio,
# precio_antes, descuento, stock, recomendado, promo_flash, promo_flash_vence,
# promo_flash_donde, descripcion, imagenes
COLUMNAS = [
    ("slug", "Código (no tocar)", GRIS, 42, False),
    ("titulo", "Nombre", AMARILLO, 48, True),
    ("marca", "Marca", AMARILLO, 18, True),
    ("presentacion", "Presentación (opcional)", AMARILLO, 24, True),
    ("categoria", "Categoría", AMARILLO, 26, True),
    ("subcategoria", "Subcategoría", AMARILLO, 24, True),
    ("precio", "Precio ($)", AMARILLO, 13, True),
    ("precio_antes", "Precio anterior ($, vacío si no hay oferta)", AMARILLO, 18, True),
    ("descuento", "% descuento (se calcula solo)", CALCULO, 14, False),
    ("stock", "¿Hay existencias? (SI/NO)", AMARILLO, 16, True),
    ("recomendado", "¿Recomendado en portada? (SI/NO)", AMARILLO, 18, True),
    ("promo_flash", "¿Promoción relámpago? (SI/NO)", AMARILLO, 18, True),
    ("promo_flash_vence", "Vence la promoción (AAAA-MM-DD HH:MM; vacío = no vence)", AMARILLO, 30, True),
    ("promo_flash_donde", "¿Dónde aplica la promoción? (ej: solo en la web)", AMARILLO, 34, True),
    ("descripcion", "Descripción (sin datos médicos)", AMARILLO, 60, True),
    ("imagenes", "Imagen (no tocar)", GRIS, 38, False),
]

SI = {"SI", "SÍ", "TRUE", "VERDADERO", "1"}
NO = {"NO", "FALSE", "FALSO", "0"}

VENCE_PLANILLA = re.compile(r"^(\d{4}-\d{2}-\d{2}) (\d{2}):(\d{2})$")
VENCE_ISO = re.compile(r"^(\d{4}-\d{2}-\d{2})T(\d{2}):(\d{2}):00-05:00$")


def vence_a_planilla(valor):
    if not valor:
        return ""
    coincide = VENCE_ISO.match(valor)
    if not coincide:
        return valor
    fecha, hora, minuto = coincide.groups()
    return f"{fecha} {hora}:{minuto}"


def vence_de_planilla(valor, quien, errores):
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None, True
    texto = str(valor).strip()
    coincide = VENCE_PLANILLA.match(texto)
    if not coincide:
        errores.append(
            f"{quien}: 'vence la promoción' debe escribirse AAAA-MM-DD HH:MM "
            f"(por ejemplo 2026-10-05 20:00, hora de Bogotá) o quedar vacío; "
            f"tiene '{valor}'"
        )
        return None, False
    fecha, hora, minuto = coincide.groups()
    return f"{fecha}T{hora}:{minuto}:00-05:00", True


def a_si_no(valor):
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return ""
    if isinstance(valor, bool):
        return "SI" if valor else "NO"
    return str(valor).strip().upper()


def de_si_no(valor, campo, quien, errores):
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return False, True
    texto = str(valor).strip().upper()
    if texto in SI:
        return True, True
    if texto in NO:
        return False, True
    errores.append(
        f"{quien}: '{campo}' debe ser SI o NO, sin nada más; tiene '{valor}'"
    )
    return False, False


def cargar_mapa():
    """Devuelve (categorias, nombre->id, sub nombre->original, sub id->nombre)."""
    categorias = json.loads(RUTA_CATEGORIAS.read_text(encoding="utf-8"))
    cat_por_nombre = {c["nombre"]: c["id"] for c in categorias}
    sub_por_nombre = {}
    sub_por_id = {}
    for c in categorias:
        for sub in c.get("subcategorias", []):
            sub_por_nombre[sub["nombre"]] = sub["original"]
            sub_por_id[sub["id"]] = sub["nombre"]
            sub_por_id[sub["original"]] = sub["nombre"]
    return categorias, cat_por_nombre, sub_por_nombre, sub_por_id


def generar():
    productos = json.loads(RUTA_JSON.read_text(encoding="utf-8"))
    categorias, _cat_por_nombre, _sub_por_nombre, sub_por_id = cargar_mapa()
    cat_por_id = {c["id"]: c["nombre"] for c in categorias}

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Productos"
    hoja.sheet_properties.tabColor = "FB01C5"
    hoja.freeze_panes = "B2"
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{len(productos) + 1}"

    for col, (campo, titulo, fondo, ancho, _editable) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=col, value=titulo)
        celda.font = FUENTE_TITULO
        celda.fill = FONDO_TITULO
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        hoja.column_dimensions[get_column_letter(col)].width = ancho
    hoja.row_dimensions[1].height = 45

    for fila, p in enumerate(productos, start=2):
        valores = {
            "slug": p.get("slug"),
            "titulo": p.get("titulo"),
            "marca": p.get("marca"),
            "presentacion": p.get("presentacion"),
            "categoria": cat_por_id.get(p.get("categoria"), p.get("categoria")),
            "subcategoria": p.get("subcategoria_nombre") or sub_por_id.get(p.get("subcategoria"), p.get("subcategoria")),
            "precio": p.get("precio"),
            "precio_antes": p.get("precio_antes"),
            "descuento": f'=IF(H{fila}="","",ROUND((H{fila}-G{fila})/H{fila}*100))',
            "stock": a_si_no(p.get("stock")),
            "recomendado": a_si_no(p.get("recomendado")),
            "promo_flash": a_si_no(p.get("promo_flash")),
            "promo_flash_vence": vence_a_planilla(p.get("promo_flash_vence")),
            "promo_flash_donde": p.get("promo_flash_donde"),
            "descripcion": p.get("descripcion"),
            "imagenes": ", ".join(p.get("imagenes") or []),
        }
        for col, (campo, _t, fondo, _a, _e) in enumerate(COLUMNAS, start=1):
            celda = hoja.cell(row=fila, column=col, value=valores[campo])
            celda.font = FUENTE
            celda.fill = fondo
            celda.alignment = Alignment(vertical="center", wrap_text=(campo == "descripcion"))
        if p.get("oculto"):
            hoja.cell(row=fila, column=1).comment = None

    # Listas para los desplegables
    listas = libro.create_sheet("Listas")
    listas.sheet_state = "hidden"
    for i, c in enumerate(categorias, start=1):
        listas.cell(row=i, column=1, value=c["nombre"])
    subs = sorted({sub["nombre"] for c in categorias for sub in c.get("subcategorias", [])})
    for i, nombre in enumerate(subs, start=1):
        listas.cell(row=i, column=2, value=nombre)
    listas.cell(row=1, column=3, value="SI")
    listas.cell(row=2, column=3, value="NO")
    n_cat = len(categorias)
    n_sub = len(subs)
    ultima = len(productos) + 1

    def desplegable(col_letra, rango_origen, mensaje, permitir_vacio=True):
        dv = DataValidation(type="list", formula1=rango_origen, allow_blank=permitir_vacio)
        dv.error = mensaje
        dv.errorTitle = "Revisar"
        dv.showErrorMessage = True
        hoja.add_data_validation(dv)
        dv.add(f"{col_letra}2:{col_letra}{ultima}")

    desplegable("E", f"Listas!$A$1:$A${n_cat}", "Elige una categoría de la lista.")
    desplegable("F", f"Listas!$B$1:$B${n_sub}", "Elige una subcategoría de la lista.")
    for col in ("J", "K", "L"):
        desplegable(col, "Listas!$C$1:$C$2", "Escribe SI o NO.")

    ent = DataValidation(type="whole", operator="greaterThanOrEqual", formula1=0, allow_blank=True)
    ent.error = "El precio es un número entero, sin puntos ni signo $ (350000). Si no lo conoces, déjalo en 0: el producto no se publica."
    ent.showErrorMessage = True
    hoja.add_data_validation(ent)
    ent.add(f"G2:G{ultima}")
    ent2 = copy(ent)
    ent2.operator = "greaterThan"
    ent2.formula1 = 0
    ent2.error = "El precio anterior es un número entero mayor que cero, sin puntos ni $ (420000). Si no hay oferta, déjalo vacío."
    hoja.add_data_validation(ent2)
    ent2.add(f"H2:H{ultima}")

    # Hoja de referencia de categorias
    ref = libro.create_sheet("Categorías")
    ref.sheet_properties.tabColor = "6FE7FF"
    ref.append(["Categoría", "Subcategoría", "Código interno"])
    for celda in ref[1]:
        celda.font = FUENTE_TITULO
        celda.fill = FONDO_TITULO
    for c in categorias:
        if not c.get("subcategorias"):
            ref.append([c["nombre"], "", c["id"]])
        for sub in c.get("subcategorias", []):
            ref.append([c["nombre"], sub["nombre"], sub["original"]])
    ref.column_dimensions["A"].width = 34
    ref.column_dimensions["B"].width = 32
    ref.column_dimensions["C"].width = 30
    ref.freeze_panes = "A2"

    instrucciones = libro.create_sheet("Instrucciones")
    instrucciones.sheet_properties.tabColor = "FFD400"
    instrucciones.column_dimensions["A"].width = 112
    lineas = [
        "CÓMO USAR ESTA PLANILLA",
        "",
        "1. Edita solo las celdas en AMARILLO. Las grises (código e imagen) no se tocan.",
        "2. El precio se escribe como número entero, sin puntos ni signo $: 350000.",
        "3. Si un producto no tiene precio todavía, déjalo en 0. Queda guardado pero no se publica.",
        "4. El % descuento se calcula solo. Para poner una oferta, llena 'Precio anterior' con un valor",
        "   mayor al precio actual. Si no hay oferta, deja el precio anterior vacío.",
        "5. 'Recomendado' marca los productos que salen en la sección Recomendaciones de la portada.",
        "6. 'Promoción relámpago' activa la ventana emergente de ofertas. Solo funciona con productos",
        "   disponibles; si el producto está agotado, no se muestra.",
        "7. 'Vence la promoción' es opcional: AAAA-MM-DD HH:MM en hora de Bogotá. Si se deja vacía,",
        "   la promoción sigue hasta que alguien la apague a mano.",
        "8. '¿Dónde aplica la promoción?' es texto libre: 'Solo en la página web', 'En todas las sedes', etc.",
        "9. Categoría y subcategoría se eligen de la lista. Si no aparece la que buscas, avísame:",
        "   no la escribas a mano ni inventes una nueva.",
        "10. No agregues ni borres filas ni columnas. Para un producto nuevo o eliminado, avísame y lo hago yo.",
        "11. Las imágenes se cambian con el sistema, no desde esta planilla. La columna queda solo de referencia.",
        "12. Guarda y envíame el archivo. Yo lo reviso producto por producto: si algo quedó mal, te digo",
        "    qué producto, qué pasa y cómo se escribe bien. Nada se publica hasta corregirlo.",
        "13. La descripción puede mejorarse (material, medidas, contenido), pero no inventes propiedades",
        "    médicas ni resultados garantizados. Es una tienda para mayores de 18 años.",
    ]
    for i, linea in enumerate(lineas, start=1):
        celda = instrucciones.cell(row=i, column=1, value=linea)
        celda.font = Font(name="Arial", size=12, bold=(i == 1))
        celda.alignment = Alignment(wrap_text=True, vertical="top")

    libro.save(RUTA_XLSX)
    print(f"Planilla generada: {RUTA_XLSX} ({len(productos)} productos, {len(categorias)} categorías)")


def leer_filas(ruta_xlsx):
    libro = load_workbook(ruta_xlsx, data_only=False)
    hoja = libro["Productos"]
    filas = []
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if not fila[0]:
            continue
        filas.append({campo: fila[i] for i, (campo, *_resto) in enumerate(COLUMNAS)})
    return filas


def construir_productos(filas, base, categorias):
    """Mezcla lo editado con el JSON actual. Devuelve (productos, errores)."""
    errores = []
    _, cat_por_nombre, sub_por_nombre, sub_por_id = cargar_mapa()
    cat_ids = {c["id"] for c in categorias}
    por_slug = {p["slug"]: p for p in base}
    vistos = set()
    productos = []

    for f in filas:
        slug = str(f["slug"]).strip()
        if slug in vistos:
            errores.append(f"{slug}: la fila está repetida en la planilla")
            continue
        vistos.add(slug)
        if slug not in por_slug:
            errores.append(
                f"{slug}: no existe en el catálogo. No agregues filas: "
                f"para un producto nuevo, avísame y lo creo yo."
            )
            continue
        original = por_slug[slug]
        p = dict(original)

        for campo in ("titulo", "marca", "presentacion", "descripcion", "promo_flash_donde"):
            v = f[campo]
            p[campo] = None if v is None or (isinstance(v, str) and not v.strip()) else str(v).strip()

        # Categoria: acepta el nombre legible o el id interno.
        cat = f["categoria"]
        if cat is None or not str(cat).strip():
            p["categoria"] = original.get("categoria")
        else:
            texto = str(cat).strip()
            p["categoria"] = cat_por_nombre.get(texto, texto if texto in cat_ids else texto)

        # Subcategoria: acepta el nombre legible, el id o el rotulo original.
        sub = f["subcategoria"]
        if sub is None or not str(sub).strip():
            p["subcategoria"] = original.get("subcategoria")
        else:
            texto = str(sub).strip()
            p["subcategoria"] = sub_por_nombre.get(texto, texto)

        for campo in ("precio", "precio_antes"):
            v = f[campo]
            if v is None or (isinstance(v, str) and not v.strip()):
                p[campo] = None if campo == "precio_antes" else 0
            elif isinstance(v, bool) or not isinstance(v, (int, float)) or int(v) != v:
                errores.append(
                    f"{slug}: '{campo}' debe ser un número entero "
                    f"(350000, sin puntos ni $); tiene '{v}'"
                )
            else:
                p[campo] = int(v)

        for campo in ("stock", "recomendado", "promo_flash"):
            p[campo], _ok = de_si_no(f[campo], campo, slug, errores)

        p["promo_flash_vence"], _ok = vence_de_planilla(f["promo_flash_vence"], slug, errores)

        # El descuento siempre sale de los precios, nunca de la planilla.
        if p.get("precio_antes") and isinstance(p.get("precio"), int) and p["precio"] > 0:
            p["descuento"] = descuento_esperado(p["precio"], p["precio_antes"])
        else:
            p["descuento"] = None
            if p.get("precio_antes") is not None and p.get("precio") == 0:
                p["precio_antes"] = None

        # Sin precio no se publica. Con precio vuelve a estar visible.
        p["oculto"] = not isinstance(p.get("precio"), int) or p["precio"] <= 0
        p["subcategoria_nombre"] = sub_por_id.get(p["subcategoria"]) or original.get("subcategoria_nombre")
        p["imagenes"] = original.get("imagenes")
        productos.append(p)

    faltantes = [s for s in por_slug if s not in vistos]
    for s in faltantes:
        errores.append(f"{s}: falta su fila en la planilla. No borres filas.")

    orden = [p["slug"] for p in base]
    productos.sort(key=lambda p: orden.index(p["slug"]))
    return productos, errores


def actualizar_conteos(productos):
    """Mantiene los totales de categorias.json al dia despues de editar la planilla."""
    visibles = [p for p in productos if not p.get("oculto")]
    categorias = json.loads(RUTA_CATEGORIAS.read_text(encoding="utf-8"))
    for grupo in categorias:
        propios = [p for p in visibles if p.get("categoria") == grupo["id"]]
        grupo["total"] = len(propios)
        for sub in grupo.get("subcategorias", []):
            sub["total"] = sum(1 for p in propios if p.get("subcategoria") == sub.get("original"))
        grupo["subcategorias"] = [s for s in grupo.get("subcategorias", []) if s["total"] > 0]
    RUTA_CATEGORIAS.write_text(
        json.dumps(categorias, ensure_ascii=False, indent=1) + "\n", encoding="utf-8"
    )


def revisar(ruta_xlsx):
    base = json.loads(RUTA_JSON.read_text(encoding="utf-8"))
    categorias = json.loads(RUTA_CATEGORIAS.read_text(encoding="utf-8"))
    filas = leer_filas(ruta_xlsx)
    productos, errores = construir_productos(filas, base, categorias)
    cat_ids, sub_ids = cargar_categorias(RUTA_JSON)
    errores += validar(productos, None, cat_ids, sub_ids)
    if errores:
        print(f"{len(errores)} problema(s) en {ruta_xlsx}:\n")
        for mensaje in errores[:60]:
            print(f"  - {mensaje}")
        if len(errores) > 60:
            print(f"  ... y {len(errores) - 60} mas")
        print("\nNada se aplicó al catálogo.")
        return 1
    print(f"La planilla está bien: {len(productos)} productos, 0 problemas.")
    return 0


def aplicar(ruta_xlsx):
    base = json.loads(RUTA_JSON.read_text(encoding="utf-8"))
    categorias = json.loads(RUTA_CATEGORIAS.read_text(encoding="utf-8"))
    filas = leer_filas(ruta_xlsx)
    productos, errores = construir_productos(filas, base, categorias)
    cat_ids, sub_ids = cargar_categorias(RUTA_JSON)
    errores += validar(productos, None, cat_ids, sub_ids)
    if errores:
        print(f"{len(errores)} problema(s), no se aplicó nada:\n")
        for mensaje in errores[:60]:
            print(f"  - {mensaje}")
        return 1
    RUTA_JSON.write_text(
        json.dumps(productos, ensure_ascii=False, indent=1) + "\n", encoding="utf-8"
    )
    actualizar_conteos(productos)
    print(f"Aplicado al catálogo: {len(productos)} productos.")
    generar()
    return 0


if __name__ == "__main__":
    if len(sys.argv) < 2 or sys.argv[1] not in ("generar", "revisar", "aplicar"):
        sys.exit(__doc__)
    accion = sys.argv[1]
    if accion == "generar":
        sys.exit(generar() or 0)
    if len(sys.argv) < 3:
        sys.exit("Falta el archivo: revisar|aplicar Catalogo-Adultoyss.xlsx")
    ruta = Path(sys.argv[2])
    if not ruta.exists():
        sys.exit(f"No existe el archivo: {ruta}")
    sys.exit(revisar(ruta) if accion == "revisar" else aplicar(ruta))
