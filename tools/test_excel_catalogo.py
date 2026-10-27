"""Prueba de extremo a extremo de excel_catalogo.py: generar -> revisar -> aplicar.

   python3 -m unittest discover -s tools -p "test_*.py" -v
"""
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import excel_catalogo as ec  # noqa: E402

CATEGORIAS = [
    {
        "id": "vibradores",
        "nombre": "Vibradores y succionadores",
        "subcategorias": [
            {"id": "balas", "nombre": "Balas vibradoras", "original": "BALAS"},
        ],
    },
    {
        "id": "lenceria",
        "nombre": "Lencería y disfraces",
        "subcategorias": [
            {"id": "lenceria", "nombre": "Lencería", "original": "LENCERIA"},
        ],
    },
]

PRODUCTOS = [
    {
        "slug": "bala-afrodita",
        "titulo": "Bala Afrodita",
        "marca": "Afrodita",
        "presentacion": "",
        "precio": 65000,
        "precio_antes": None,
        "descuento": None,
        "categoria": "vibradores",
        "subcategoria": "BALAS",
        "subcategoria_nombre": "Balas vibradoras",
        "stock": True,
        "recomendado": False,
        "promo_flash": False,
        "promo_flash_vence": None,
        "promo_flash_donde": None,
        "imagenes": ["productos/bala-afrodita.webp"],
        "descripcion": "Bala recargable.",
        "oculto": False,
    },
    {
        "slug": "producto-sin-precio",
        "titulo": "Producto sin precio",
        "marca": "",
        "presentacion": "",
        "precio": 0,
        "precio_antes": None,
        "descuento": None,
        "categoria": "lenceria",
        "subcategoria": "LENCERIA",
        "subcategoria_nombre": "Lencería",
        "stock": True,
        "recomendado": False,
        "promo_flash": False,
        "promo_flash_vence": None,
        "promo_flash_donde": None,
        "imagenes": ["productos/producto-sin-precio.webp"],
        "descripcion": "Pendiente de precio.",
        "oculto": True,
    },
]


class ExcelCatalogoTest(unittest.TestCase):
    def setUp(self):
        self.temporal = tempfile.TemporaryDirectory()
        raiz = Path(self.temporal.name)
        (raiz / "productos.json").write_text(
            json.dumps(PRODUCTOS, ensure_ascii=False, indent=1), encoding="utf-8"
        )
        (raiz / "categorias.json").write_text(
            json.dumps(CATEGORIAS, ensure_ascii=False, indent=1), encoding="utf-8"
        )
        self.rutas = (ec.RUTA_JSON, ec.RUTA_CATEGORIAS, ec.RUTA_XLSX)
        ec.RUTA_JSON = raiz / "productos.json"
        ec.RUTA_CATEGORIAS = raiz / "categorias.json"
        ec.RUTA_XLSX = raiz / "Catalogo-Adultoyss.xlsx"

    def tearDown(self):
        ec.RUTA_JSON, ec.RUTA_CATEGORIAS, ec.RUTA_XLSX = self.rutas
        self.temporal.cleanup()

    def fila_de(self, hoja, slug):
        for fila in range(2, hoja.max_row + 1):
            if hoja.cell(fila, 1).value == slug:
                return fila
        raise AssertionError(f"no se encontro {slug}")

    def test_generar_y_revisar_sin_cambios(self):
        ec.generar()
        self.assertTrue(ec.RUTA_XLSX.exists())
        self.assertEqual(ec.revisar(ec.RUTA_XLSX), 0)

    def test_aplicar_precio_descuento_y_recomendado(self):
        ec.generar()
        libro = load_workbook(ec.RUTA_XLSX)
        hoja = libro["Productos"]
        fila = self.fila_de(hoja, "bala-afrodita")
        hoja.cell(fila, 7).value = 60000
        hoja.cell(fila, 8).value = 75000
        hoja.cell(fila, 11).value = "SI"
        libro.save(ec.RUTA_XLSX)

        self.assertEqual(ec.aplicar(ec.RUTA_XLSX), 0)
        productos = {p["slug"]: p for p in json.loads(ec.RUTA_JSON.read_text(encoding="utf-8"))}
        p = productos["bala-afrodita"]
        self.assertEqual(p["precio"], 60000)
        self.assertEqual(p["precio_antes"], 75000)
        self.assertEqual(p["descuento"], 20)
        self.assertTrue(p["recomendado"])
        self.assertFalse(p["oculto"])

    def test_vence_acepta_celda_de_fecha_de_excel(self):
        # Al escribir 2026-11-30 22:00 en Excel, la celda queda como fecha real.
        from datetime import datetime
        ec.generar()
        libro = load_workbook(ec.RUTA_XLSX)
        hoja = libro["Productos"]
        fila = self.fila_de(hoja, "bala-afrodita")
        hoja.cell(fila, 12).value = "SI"
        hoja.cell(fila, 13).value = datetime(2026, 11, 30, 22, 0)
        libro.save(ec.RUTA_XLSX)

        self.assertEqual(ec.aplicar(ec.RUTA_XLSX), 0)
        productos = {p["slug"]: p for p in json.loads(ec.RUTA_JSON.read_text(encoding="utf-8"))}
        self.assertEqual(productos["bala-afrodita"]["promo_flash_vence"], "2026-11-30T22:00:00-05:00")

    def test_precio_cero_oculta_el_producto(self):
        ec.generar()
        libro = load_workbook(ec.RUTA_XLSX)
        hoja = libro["Productos"]
        fila = self.fila_de(hoja, "bala-afrodita")
        hoja.cell(fila, 7).value = 0
        libro.save(ec.RUTA_XLSX)

        self.assertEqual(ec.aplicar(ec.RUTA_XLSX), 0)
        p = {x["slug"]: x for x in json.loads(ec.RUTA_JSON.read_text(encoding="utf-8"))}["bala-afrodita"]
        self.assertTrue(p["oculto"])
        self.assertIsNone(p["descuento"])

    def test_si_no_mal_escrito_no_aplica_nada(self):
        ec.generar()
        antes = ec.RUTA_JSON.read_text(encoding="utf-8")
        libro = load_workbook(ec.RUTA_XLSX)
        hoja = libro["Productos"]
        fila = self.fila_de(hoja, "bala-afrodita")
        hoja.cell(fila, 10).value = "TAL VEZ"
        libro.save(ec.RUTA_XLSX)

        self.assertEqual(ec.revisar(ec.RUTA_XLSX), 1)
        self.assertEqual(ec.aplicar(ec.RUTA_XLSX), 1)
        self.assertEqual(ec.RUTA_JSON.read_text(encoding="utf-8"), antes)

    def test_categoria_y_subcategoria_por_nombre(self):
        ec.generar()
        libro = load_workbook(ec.RUTA_XLSX)
        hoja = libro["Productos"]
        fila = self.fila_de(hoja, "bala-afrodita")
        hoja.cell(fila, 5).value = "Lencería y disfraces"
        hoja.cell(fila, 6).value = "Lencería"
        libro.save(ec.RUTA_XLSX)

        self.assertEqual(ec.aplicar(ec.RUTA_XLSX), 0)
        p = {x["slug"]: x for x in json.loads(ec.RUTA_JSON.read_text(encoding="utf-8"))}["bala-afrodita"]
        self.assertEqual(p["categoria"], "lenceria")
        self.assertEqual(p["subcategoria"], "LENCERIA")
        self.assertEqual(p["subcategoria_nombre"], "Lencería")


if __name__ == "__main__":
    unittest.main()
